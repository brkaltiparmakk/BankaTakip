"""Excel (.xls/.xlsx) ve HTML tablo eklerini satır satır metne çevirir.

Bazı bankalar (ör. Garanti BBVA "Hesap Hareketleri") hareketleri PDF yerine Excel
dosyası olarak gönderir; bu dosyalar bazen gerçekte Excel görünümlü bir HTML tablosudur.
Çıkan metin PDF'lerdeki gibi GenericParser ile ayrıştırılır: her satır hücreleri
aralarında iki boşlukla birleştirilmiş tek bir satır olur.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from html.parser import HTMLParser

XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # eski Excel (OLE2)
ZIP_MAGIC = b"PK\x03\x04"                          # xlsx


def detect_kind(content: bytes) -> str | None:
    if content.startswith(b"%PDF"):
        return "pdf"
    if content.startswith(XLS_MAGIC):
        return "xls"
    if content.startswith(ZIP_MAGIC):
        return "xlsx"
    head = content[:2048].lower()
    if b"<table" in head or b"<html" in head or b"<?xml" in head and b"<table" in content[:20000].lower():
        return "html"
    return None


def _fmt_number(value: float) -> str:
    """Excel'deki sayıyı Türkçe biçime çevirir: -1234.5 → -1.234,50"""
    text = f"{abs(value):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"-{text}" if value < 0 else text


def _fmt_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            return _fmt_number(value)
        # Tam sayılar: tutar olabilir (1500 → 1.500,00) ama dekont/hesap no da olabilir.
        # Tutar sütunlarını ondalıklı yazmak için tam sayıyı da ondalıklı biçimliyoruz;
        # 8 haneden uzun sayılar kimlik numarası kabul edilir.
        return str(int(value)) if abs(value) >= 10**8 else _fmt_number(float(value))
    return " ".join(str(value).split())


def rows_from_xlsx(content: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            rows.append([_fmt_cell(v) for v in row])
    return rows


def rows_from_xls(content: bytes) -> list[list[str]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    rows = []
    for sheet in book.sheets():
        for r in range(sheet.nrows):
            cells = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    cells.append(_fmt_cell(xlrd.xldate_as_datetime(cell.value, book.datemode)))
                else:
                    cells.append(_fmt_cell(cell.value))
            rows.append(cells)
    return rows


class _TableParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows: list[list[str]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _decode_html(content: bytes) -> str:
    head = content[:4096].lower()
    for enc in ("utf-8", "windows-1254", "iso-8859-9"):
        if enc.encode() in head:
            try:
                return content.decode(enc)
            except UnicodeDecodeError:
                break
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("windows-1254", errors="replace")


def rows_from_html(content: bytes) -> list[list[str]]:
    parser = _TableParser()
    parser.feed(_decode_html(content))
    return parser.rows


def rows_to_text(rows: list[list[str]]) -> str:
    lines = []
    for row in rows:
        cells = [c for c in row if c]
        if cells:
            lines.append("  ".join(cells))
    return "\n".join(lines)


def extract_table_text(content: bytes, kind: str) -> str:
    readers = {"xlsx": rows_from_xlsx, "xls": rows_from_xls, "html": rows_from_html}
    return rows_to_text(readers[kind](content))
